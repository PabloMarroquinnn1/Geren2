"""Carga completa de datos de QuetzalMart en Odoo (desde cualquier computadora).

Usa la API JSON-RPC de Odoo con un usuario administrador. Las plantillas de Excel
se importan con el mismo motor que el botón "Importar registros" de la interfaz.

Uso:
    python carga_masiva/cargar_quetzalmart.py --url https://MI-DOMINIO --clave CLAVE_ADMIN
    python carga_masiva/cargar_quetzalmart.py ... --paso transacciones   (solo un paso)

Pasos (en orden): configurar, maestros, imagenes, existencias, crm, transacciones,
documentos, verificar. "todo" ejecuta todos. Cada paso puede repetirse sin duplicar.
"""
import argparse
import base64
import csv
import io
import random
import sys
import time
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook

import catalogo
from odoo_rpc import ErrorOdoo, Odoo

AQUI = Path(__file__).resolve().parent
RAIZ = AQUI.parent
PLANTILLAS = AQUI / "plantillas"
IMAGENES = AQUI / "imagenes" / "productos"
sys.path.insert(0, str(RAIZ / "gestion_documental"))

ORDEN_PLANTILLAS = [
    ("01_departamentos", "hr.department"),
    ("02_puestos", "hr.job"),
    ("03_empleados", "hr.employee"),
    ("04_contratos", "hr.contract"),
    ("05_clientes", "res.partner"),
    ("06_proveedores", "res.partner"),
    ("07_productos", "product.template"),
    ("08_materiales_operativos", "product.template"),
]
DESTACADOS = ["DES-009", "LAC-001", "BEB-003", "SNK-004", "LIM-001", "CUI-001", "DES-004", "LAC-007"]
SUCURSAL_POR_PAIS = {"GT": "GUA", "MX": "MEX", "SV": "SAL"}

azar = random.Random(150)


def titulo(texto):
    print("\n\033[1;32m== %s ==\033[0m" % texto)


# ---------------------------------------------------------------------------
# 1. Configuración del sitio y del correo saliente
# ---------------------------------------------------------------------------
def configurar(odoo, args):
    titulo("Configuración general")
    parametros = {"web.base.url": odoo.url, "web.base.url.freeze": "True", "dms.binary_max_size": "50"}
    for clave, valor in parametros.items():
        odoo.ejecutar("ir.config_parameter", "set_param", clave, valor)
    sitio = odoo.buscar("website", [], limit=1)
    valores_sitio = {"domain": odoo.url}
    if args.ga4:
        valores_sitio["google_analytics_key"] = args.ga4
        print("  Google Analytics 4:", args.ga4)
    odoo.escribir("website", sitio, valores_sitio)
    if args.smtp_servidor:
        servidor = {
            "name": "Correo QuetzalMart",
            "smtp_host": args.smtp_servidor,
            "smtp_port": args.smtp_puerto,
            "smtp_encryption": {587: "starttls", 465: "ssl"}.get(args.smtp_puerto, "none"),
            "smtp_authentication": "login",
            "smtp_user": args.smtp_usuario,
            "smtp_pass": args.smtp_clave,
            "from_filter": args.correo_remitente,
            "sequence": 1,
        }
        existentes = odoo.buscar("ir.mail_server", [("name", "=", "Correo QuetzalMart")])
        if existentes:
            odoo.escribir("ir.mail_server", existentes, servidor)
        else:
            odoo.crear("ir.mail_server", servidor)
        print("  Servidor SMTP:", args.smtp_servidor, "como", args.correo_remitente)
    if args.correo_remitente:
        empresa = odoo.ref("base.main_company")
        odoo.escribir("res.company", [empresa], {"email": args.correo_remitente})
        odoo.ejecutar("ir.config_parameter", "set_param", "mail.default.from", args.correo_remitente)
        odoo.ejecutar("ir.config_parameter", "set_param", "mail.catchall.domain",
                      args.correo_remitente.split("@")[1])
        print("  Remitente de los correos:", args.correo_remitente)


# ---------------------------------------------------------------------------
# 2. Datos maestros con el importador de Odoo
# ---------------------------------------------------------------------------
def importar_plantilla(odoo, nombre, modelo):
    """Importa la primera hoja de la plantilla con el asistente "Importar registros".

    El asistente guarda el archivo en bruto (no en base64); por JSON-RPC solo se
    puede enviar texto, así que la hoja viaja como CSV con el mismo contenido.
    """
    ruta = PLANTILLAS / f"{nombre}.xlsx"
    hoja = load_workbook(ruta, read_only=True).worksheets[0]
    filas = [["" if valor is None else valor for valor in fila]
             for fila in hoja.iter_rows(values_only=True)]
    encabezados = filas[0]
    texto = io.StringIO()
    csv.writer(texto).writerows(filas)
    asistente = odoo.crear("base_import.import", {
        "res_model": modelo,
        "file": texto.getvalue(),
        "file_name": ruta.with_suffix(".csv").name,
        "file_type": "text/csv",
    })
    opciones = {
        "has_headers": True, "advanced": True, "keep_matches": False,
        "date_format": "%Y-%m-%d", "datetime_format": "", "float_thousand_separator": ",",
        "float_decimal_separator": ".", "separator": ",", "quoting": '"', "encoding": "utf-8",
        "name_create_enabled_fields": {"category_id": True}, "import_set_empty_fields": [],
        "import_skip_records": [], "fallback_values": {}, "skip": 0, "limit": 5000,
        "tracking_disable": True,
    }
    vista = odoo.ejecutar("base_import.import", "parse_preview", [asistente], opciones)
    if vista.get("error"):
        raise ErrorOdoo("%s: %s" % (nombre, vista["error"]))
    resultado = odoo.ejecutar("base_import.import", "execute_import", [asistente],
                              encabezados, encabezados, opciones)
    errores = [m for m in resultado.get("messages", []) if m.get("type") == "error"]
    if errores:
        for mensaje in errores[:5]:
            print("   ", mensaje.get("message"), mensaje.get("rows", ""))
        raise ErrorOdoo("La plantilla %s tiene errores" % nombre)
    print(f"  {nombre:<30} {len(resultado.get('ids') or []):>4} registros -> {modelo}")


def maestros(odoo, args):
    titulo("Datos maestros (importación de plantillas)")
    for nombre, modelo in ORDEN_PLANTILLAS:
        importar_plantilla(odoo, nombre, modelo)
    # Lista de correo de la campaña masiva con todos los clientes.
    lista = odoo.ref("qm_crm_postventa.lista_clientes")
    clientes = odoo.leer("res.partner", [("customer_rank", ">", 0), ("email", "!=", False)],
                         ["name", "email", "country_id"])
    existentes = {c["email"] for c in odoo.leer("mailing.contact", [], ["email"])}
    nuevos = [{"name": c["name"], "email": c["email"],
               "country_id": c["country_id"] and c["country_id"][0],
               "list_ids": [(4, lista)]} for c in clientes if c["email"] not in existentes]
    if nuevos:
        odoo.ejecutar("mailing.contact", "create", nuevos)
    print(f"  Lista 'Clientes QuetzalMart': {len(clientes)} contactos")


def imagenes(odoo, args):
    titulo("Imágenes de productos y destacados de la tienda")
    productos = odoo.leer("product.template", [("default_code", "!=", False)], ["default_code"])
    cargadas = 0
    for producto in productos:
        archivo = IMAGENES / f"{producto['default_code']}.png"
        if archivo.exists():
            odoo.escribir("product.template", [producto["id"]],
                          {"image_1920": base64.b64encode(archivo.read_bytes()).decode()})
            cargadas += 1
    for orden, sku in enumerate(DESTACADOS, 1):
        odoo.escribir("product.template", odoo.buscar("product.template", [("default_code", "=", sku)]),
                      {"website_sequence": orden})
    print(f"  {cargadas} imágenes cargadas; {len(DESTACADOS)} productos destacados")


# ---------------------------------------------------------------------------
# 3. Existencias iniciales por sucursal y sector
# ---------------------------------------------------------------------------
def existencias(odoo, args):
    titulo("Existencias iniciales (inventario físico)")
    if odoo.contar("stock.move", [("is_inventory", "=", True), ("state", "=", "done")]) > 100:
        print("  Ya existen ajustes de inventario: se omite.")
        return
    libro = load_workbook(PLANTILLAS / "09_existencias_iniciales.xlsx", read_only=True)
    filas = list(libro.worksheets[0].iter_rows(min_row=2, values_only=True))
    productos = {p["default_code"]: p["id"] for p in odoo.leer(
        "product.product", [("default_code", "!=", False)], ["default_code"])}
    ubicaciones = {}
    quants = []
    for sku, ubicacion_xmlid, cantidad in filas:
        if ubicacion_xmlid not in ubicaciones:
            ubicaciones[ubicacion_xmlid] = odoo.ref(ubicacion_xmlid)
        quants.append({"product_id": productos[sku], "location_id": ubicaciones[ubicacion_xmlid],
                       "inventory_quantity": cantidad})
    for inicio in range(0, len(quants), 100):
        ids = odoo.ejecutar("stock.quant", "create", quants[inicio:inicio + 100],
                            context={"inventory_mode": True})
        odoo.ejecutar("stock.quant", "action_apply_inventory", ids,
                      context={"inventory_mode": True, "inventory_name": "Existencias iniciales QuetzalMart"})
    print(f"  {len(quants)} existencias registradas en {len(ubicaciones)} ubicaciones")


def crm(odoo, args):
    titulo("Oportunidades del CRM")
    importar_plantilla(odoo, "10_oportunidades_crm", "crm.lead")


# ---------------------------------------------------------------------------
# 4. Transacciones: cotizaciones, compras, ventas, facturas y pagos
# ---------------------------------------------------------------------------
def fecha_aleatoria(inicio, fin):
    dias = (fin - inicio).days
    # Más movimiento en los meses recientes (la empresa está creciendo).
    return inicio + timedelta(days=int(dias * (azar.random() ** 0.7)))


def a_texto(fecha, hora=10):
    return datetime(fecha.year, fecha.month, fecha.day, hora + azar.randint(0, 8),
                    azar.randint(0, 59)).strftime("%Y-%m-%d %H:%M:%S")


def validar_albaran(odoo, picking_ids):
    for picking in picking_ids:
        estado = odoo.leer("stock.picking", [("id", "=", picking)], ["state"])[0]["state"]
        if estado in ("done", "cancel"):
            continue
        odoo.ejecutar("stock.picking", "action_assign", [picking])
        accion = odoo.ejecutar("stock.picking", "button_validate", [picking])
        if isinstance(accion, dict) and accion.get("res_model") == "stock.backorder.confirmation":
            asistente = odoo.crear("stock.backorder.confirmation", {
                "pick_ids": [(6, 0, [picking])]}, contexto=accion.get("context"))
            odoo.ejecutar("stock.backorder.confirmation", "process_cancel_backorder", [asistente],
                          context=accion.get("context"))


def pagar(odoo, factura_id, fecha):
    contexto = {"active_model": "account.move", "active_ids": [factura_id]}
    asistente = odoo.crear("account.payment.register", {"payment_date": fecha.isoformat()}, contexto)
    odoo.ejecutar("account.payment.register", "action_create_payments", [asistente], context=contexto)


def datos_para_transacciones(odoo):
    almacenes = {}
    for codigo in ("gua", "mex", "sal"):
        almacen = odoo.leer("stock.warehouse", [("id", "=", odoo.ref(f"qm_implementacion.almacen_{codigo}"))],
                            ["in_type_id"])[0]
        almacenes[codigo.upper()] = {"id": almacen["id"], "entrada": almacen["in_type_id"][0],
                                     "equipo": odoo.ref(f"qm_implementacion.equipo_{codigo}")}
    clientes = odoo.leer("res.partner", [("customer_rank", ">", 0), ("ref", "like", "CLI-")],
                         ["country_id", "is_company"])
    productos = odoo.leer("product.product", [("default_code", "!=", False), ("sale_ok", "=", True)],
                          ["default_code"])
    materiales = odoo.leer("product.product", [("default_code", "like", "MAT-")], ["default_code"])
    proveedores = odoo.leer("res.partner", [("supplier_rank", ">", 0), ("ref", "like", "PRV-")], ["ref"])
    return almacenes, clientes, productos, materiales, proveedores


def lineas_venta(productos, empresa, cantidad_lineas):
    # Algunos productos se venden más que otros: sirve para el análisis de "más vendidos".
    pesos = [6 if p["default_code"] in DESTACADOS else (3 if i % 4 == 0 else 1)
             for i, p in enumerate(productos)]
    elegidos = []
    while len(elegidos) < cantidad_lineas:
        producto = azar.choices(productos, weights=pesos)[0]
        if producto not in elegidos:
            elegidos.append(producto)
    return [(0, 0, {"product_id": p["id"],
                    "product_uom_qty": azar.randint(6, 24) if empresa else azar.randint(1, 5)})
            for p in elegidos]


def proveedor_del_producto(odoo, producto_id, cache={}):
    if producto_id not in cache:
        plantilla = odoo.leer("product.product", [("id", "=", producto_id)], ["seller_ids"])[0]
        vendedor = odoo.leer("product.supplierinfo", [("id", "in", plantilla["seller_ids"])], ["partner_id"],
                             limit=1)
        cache[producto_id] = vendedor[0]["partner_id"][0] if vendedor else None
    return cache[producto_id]


def transacciones(odoo, args):
    titulo("Transacciones")
    almacenes, clientes, productos, materiales, proveedores = datos_para_transacciones(odoo)
    inicio, fin = date(2026, 6, 1), date(2026, 10, 15)

    # --- 20 cotizaciones de venta (borrador / enviadas) ---------------------
    hechas = odoo.contar("sale.order", [("client_order_ref", "like", "COT-")])
    for numero in range(hechas + 1, 21):
        cliente = azar.choice(clientes)
        sucursal = almacenes[SUCURSAL_POR_PAIS.get(_codigo_pais(odoo, cliente), "GUA")]
        pedido = odoo.crear("sale.order", {
            "partner_id": cliente["id"], "warehouse_id": sucursal["id"], "team_id": sucursal["equipo"],
            "client_order_ref": f"COT-{numero:03d}",
            "validity_date": (date(2026, 10, 25) + timedelta(days=numero)).isoformat(),
            "order_line": lineas_venta(productos, cliente["is_company"], azar.randint(2, 4)),
        })
        if numero % 2 == 0:
            odoo.ejecutar("sale.order", "action_quotation_sent", [pedido])
    print("  20 cotizaciones de venta")

    # --- 20 solicitudes de cotización a proveedores -------------------------
    hechas = odoo.contar("purchase.order", [("partner_ref", "like", "SDC-")])
    for numero in range(hechas + 1, 21):
        proveedor = azar.choice(proveedores)
        sucursal = almacenes[azar.choice(["GUA", "GUA", "MEX", "SAL"])]
        lineas = _lineas_compra(odoo, proveedor["id"], productos + materiales, azar.randint(2, 4))
        orden = odoo.crear("purchase.order", {
            "partner_id": proveedor["id"], "partner_ref": f"SDC-{numero:03d}",
            "picking_type_id": sucursal["entrada"], "order_line": lineas,
        })
        if numero % 2 == 0:
            odoo.escribir("purchase.order", [orden], {"state": "sent"})
    print("  20 solicitudes de cotización de compra")

    # --- 100 compras confirmadas, recibidas y facturadas ---------------------
    hechas = odoo.contar("purchase.order", [("partner_ref", "like", "OC-")])
    for numero in range(hechas + 1, 101):
        fecha = fecha_aleatoria(inicio, fin)
        proveedor = proveedores[(numero - 1) % len(proveedores)]
        sucursal = almacenes[("GUA", "GUA", "MEX", "SAL")[numero % 4]]
        lineas = _lineas_compra(odoo, proveedor["id"], productos + materiales, azar.randint(2, 5))
        orden = odoo.crear("purchase.order", {
            "partner_id": proveedor["id"], "partner_ref": f"OC-{numero:03d}",
            "picking_type_id": sucursal["entrada"], "date_order": a_texto(fecha), "order_line": lineas,
        })
        odoo.ejecutar("purchase.order", "button_confirm", [orden])
        datos = odoo.leer("purchase.order", [("id", "=", orden)], ["picking_ids"])[0]
        validar_albaran(odoo, datos["picking_ids"])
        odoo.ejecutar("purchase.order", "action_create_invoice", [orden])
        factura = odoo.leer("purchase.order", [("id", "=", orden)], ["invoice_ids"])[0]["invoice_ids"][0]
        odoo.escribir("account.move", [factura], {
            "invoice_date": (fecha + timedelta(days=2)).isoformat(),
            "ref": f"FEL-{proveedor['ref']}-{numero:05d}",
        })
        odoo.ejecutar("account.move", "action_post", [factura])
        if numero % 10 < 7:
            pagar(odoo, factura, fecha + timedelta(days=15))
        odoo.escribir("purchase.order", [orden], {"date_approve": a_texto(fecha)})
        if numero % 10 == 0:
            print(f"  compras: {numero}/100")

    # --- 150 ventas confirmadas, entregadas y facturadas ---------------------
    hechas = odoo.contar("sale.order", [("client_order_ref", "like", "VTA-")])
    for numero in range(hechas + 1, 151):
        fecha = fecha_aleatoria(inicio, fin)
        cliente = clientes[(numero * 7) % len(clientes)]
        sucursal = almacenes[SUCURSAL_POR_PAIS.get(_codigo_pais(odoo, cliente), "GUA")]
        pedido = odoo.crear("sale.order", {
            "partner_id": cliente["id"], "warehouse_id": sucursal["id"], "team_id": sucursal["equipo"],
            "client_order_ref": f"VTA-{numero:03d}",
            "order_line": lineas_venta(productos, cliente["is_company"], azar.randint(1, 5)),
        })
        odoo.ejecutar("sale.order", "action_confirm", [pedido])
        odoo.escribir("sale.order", [pedido], {"date_order": a_texto(fecha)})
        datos = odoo.leer("sale.order", [("id", "=", pedido)], ["picking_ids"])[0]
        validar_albaran(odoo, datos["picking_ids"])
        contexto = {"active_model": "sale.order", "active_ids": [pedido], "active_id": pedido}
        asistente = odoo.crear("sale.advance.payment.inv", {"advance_payment_method": "delivered"}, contexto)
        odoo.ejecutar("sale.advance.payment.inv", "create_invoices", [asistente], context=contexto)
        factura = odoo.leer("sale.order", [("id", "=", pedido)], ["invoice_ids"])[0]["invoice_ids"][0]
        odoo.escribir("account.move", [factura], {"invoice_date": fecha.isoformat()})
        odoo.ejecutar("account.move", "action_post", [factura])
        if numero % 4 != 0:
            pagar(odoo, factura, fecha + timedelta(days=azar.randint(0, 20)))
        if numero % 10 == 0:
            print(f"  ventas: {numero}/150")


def _codigo_pais(odoo, partner, cache={}):
    pais_id = partner["country_id"] and partner["country_id"][0]
    if pais_id not in cache:
        cache[pais_id] = odoo.leer("res.country", [("id", "=", pais_id)], ["code"])[0]["code"] if pais_id else "GT"
    return cache[pais_id]


def _lineas_compra(odoo, proveedor_id, candidatos, cantidad):
    propios = [p for p in candidatos if proveedor_del_producto(odoo, p["id"]) == proveedor_id]
    elegidos = azar.sample(propios, min(cantidad, len(propios))) if propios else azar.sample(candidatos, cantidad)
    es_material = any(p["default_code"].startswith("MAT-") for p in elegidos)
    return [(0, 0, {"product_id": p["id"],
                    "product_qty": azar.randint(2, 10) if es_material else azar.randint(24, 96)})
            for p in elegidos]


# ---------------------------------------------------------------------------
# 5. Gestión documental: 5 facturas de proveedor, 5 contratos de outsourcing,
#    5 contratos laborales (PDF generados con los datos del ERP)
# ---------------------------------------------------------------------------
def documentos(odoo, args):
    import generar_documentos as gd

    titulo("Gestión documental")
    pendientes = odoo.ejecutar("account.move", "qm_reintentar_archivo")
    print(f"  {pendientes} facturas en cola para archivar su PDF en Documentos")
    carpeta = RAIZ / "gestion_documental" / "documentos"
    carpeta.mkdir(exist_ok=True)
    ref = odoo.ref
    etiquetas = {nombre: ref(f"qm_gestion_documental.etiqueta_{nombre}") for nombre in
                 ("proveedor", "outsourcing", "laboral", "confidencial", "vigente", "pagada")}

    def subir(ruta, carpeta_xmlid, categoria_xmlid, tags, modelo=None, registro=None):
        directorio = ref(carpeta_xmlid)
        valores = {
            "name": ruta.name, "directory_id": directorio, "category_id": ref(categoria_xmlid),
            "tag_ids": [(6, 0, tags)], "content": base64.b64encode(ruta.read_bytes()).decode(),
            "mimetype": "application/pdf",
        }
        existente = odoo.buscar("dms.file", [("name", "=", ruta.name), ("directory_id", "=", directorio)])
        if existente:
            odoo.escribir("dms.file", existente, valores)
        else:
            odoo.crear("dms.file", valores)
        if modelo:
            if not odoo.contar("ir.attachment", [("res_model", "=", modelo), ("res_id", "=", registro),
                                                 ("name", "=", ruta.name)]):
                odoo.crear("ir.attachment", {"name": ruta.name, "res_model": modelo, "res_id": registro,
                                             "datas": valores["content"], "mimetype": "application/pdf"})

    # Facturas de proveedor: las de las 5 primeras compras confirmadas.
    facturas = odoo.leer("account.move", [("move_type", "=", "in_invoice"), ("state", "=", "posted"),
                                          ("ref", "like", "FEL-")],
                         ["name", "ref", "partner_id", "invoice_date", "invoice_line_ids", "amount_untaxed",
                          "amount_tax", "amount_total", "payment_state", "invoice_origin"],
                         order="id", limit=5)
    for factura in facturas:
        proveedor = odoo.leer("res.partner", [("id", "=", factura["partner_id"][0])],
                              ["name", "vat", "street", "city", "country_id", "phone", "email"])[0]
        lineas = odoo.leer("account.move.line", [("id", "in", factura["invoice_line_ids"]),
                                                 ("display_type", "=", "product")],
                           ["name", "quantity", "price_unit", "price_subtotal"])
        ruta = gd.factura_proveedor(carpeta, factura, proveedor, lineas)
        tags = [etiquetas["proveedor"]] + ([etiquetas["pagada"]] if factura["payment_state"] in ("paid", "in_payment") else [])
        subir(ruta, "qm_gestion_documental.carpeta_facturas_proveedor",
              "qm_gestion_documental.categoria_factura_proveedor", tags, "account.move", factura["id"])
        print("  ", ruta.name)

    # Contratos de outsourcing con los proveedores de servicios.
    for numero, (nombre, servicio, monto) in enumerate(catalogo.SERVICIOS_OUTSOURCING, 1):
        socio = odoo.leer("res.partner", [("ref", "=", f"OUT-{numero:03d}")], ["name", "vat", "street", "city"])[0]
        ruta = gd.contrato_outsourcing(carpeta, numero, socio, servicio, monto)
        subir(ruta, "qm_gestion_documental.carpeta_contratos_outsourcing",
              "qm_gestion_documental.categoria_contrato_outsourcing",
              [etiquetas["outsourcing"], etiquetas["vigente"], etiquetas["confidencial"]],
              "res.partner", socio["id"])
        print("  ", ruta.name)

    # Contratos laborales: gerentes de sucursal y personal clave.
    contratos = odoo.leer("hr.contract", [("state", "=", "open")],
                          ["name", "employee_id", "job_id", "department_id", "date_start", "wage"],
                          order="wage desc", limit=5)
    for contrato in contratos:
        empleado = odoo.leer("hr.employee", [("id", "=", contrato["employee_id"][0])],
                             ["name", "identification_id", "work_location_id", "work_email"])[0]
        ruta = gd.contrato_laboral(carpeta, contrato, empleado)
        subir(ruta, "qm_gestion_documental.carpeta_contratos_empleados",
              "qm_gestion_documental.categoria_contrato_laboral",
              [etiquetas["laboral"], etiquetas["vigente"], etiquetas["confidencial"]],
              "hr.employee", empleado["id"])
        print("  ", ruta.name)


# ---------------------------------------------------------------------------
# 6. Verificación (mismos conteos que las consultas SQL de la calificación)
# ---------------------------------------------------------------------------
def verificar(odoo, args):
    titulo("Verificación")
    conteos = [
        ("Ventas confirmadas", "sale.order", [("state", "=", "sale")], 150),
        ("Cotizaciones de venta", "sale.order", [("state", "in", ("draft", "sent"))], 20),
        ("Cotizaciones de compra", "purchase.order", [("state", "in", ("draft", "sent"))], 20),
        ("Compras confirmadas", "purchase.order", [("state", "in", ("purchase", "done"))], 100),
        ("Facturas de proveedor publicadas", "account.move",
         [("move_type", "=", "in_invoice"), ("state", "=", "posted")], 100),
        ("Facturas de cliente publicadas", "account.move",
         [("move_type", "=", "out_invoice"), ("state", "=", "posted")], 50),
        ("Empleados", "hr.employee", [], 35),
        ("Puestos", "hr.job", [], 6),
        ("Departamentos", "hr.department", [], 5),
        ("Materiales operativos", "product.product", [("default_code", "like", "MAT-")], 60),
        ("Productos publicados en la web", "product.template", [("is_published", "=", True)], 60),
        ("PDF de facturas en Documentos", "dms.file",
         [("directory_id", "=", odoo.ref("qm_gestion_documental.carpeta_facturas_venta"))], 50),
        ("Facturas de proveedor en Documentos", "dms.file",
         [("directory_id", "=", odoo.ref("qm_gestion_documental.carpeta_facturas_proveedor"))], 5),
        ("Contratos de outsourcing en Documentos", "dms.file",
         [("directory_id", "=", odoo.ref("qm_gestion_documental.carpeta_contratos_outsourcing"))], 5),
        ("Contratos de empleados en Documentos", "dms.file",
         [("directory_id", "=", odoo.ref("qm_gestion_documental.carpeta_contratos_empleados"))], 5),
    ]
    todo_bien = True
    for etiqueta, modelo, dominio, minimo in conteos:
        cantidad = odoo.contar(modelo, dominio)
        estado = "OK " if cantidad >= minimo else "FALTA"
        todo_bien &= cantidad >= minimo
        print(f"  [{estado}] {etiqueta:<40} {cantidad:>5}  (mínimo {minimo})")
    pendientes = odoo.contar("account.move", [("qm_archivo_pendiente", "=", True)])
    if pendientes:
        print(f"  Aún se están archivando {pendientes} PDF de facturas en segundo plano.")
    return todo_bien


PASOS = {
    "configurar": configurar, "maestros": maestros, "imagenes": imagenes, "existencias": existencias,
    "crm": crm, "transacciones": transacciones, "documentos": documentos, "verificar": verificar,
}


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--url", required=True, help="https://quetzalmart.<ip>.sslip.io")
    parser.add_argument("--base", default="quetzalmart")
    parser.add_argument("--usuario", default="admin@quetzalmart.com")
    parser.add_argument("--clave", required=True)
    parser.add_argument("--paso", default="todo", choices=["todo"] + list(PASOS))
    parser.add_argument("--ga4", help="ID de medición de Google Analytics 4 (G-XXXXXXXXXX)")
    parser.add_argument("--correo-remitente", help="Dirección desde la que Odoo envía los correos")
    parser.add_argument("--smtp-servidor", help="smtp.gmail.com, smtp-relay.brevo.com, ...")
    parser.add_argument("--smtp-puerto", type=int, default=587)
    parser.add_argument("--smtp-usuario")
    parser.add_argument("--smtp-clave")
    parser.add_argument("--sin-verificar-ssl", action="store_true", help="Solo para pruebas locales")
    args = parser.parse_args()

    odoo = Odoo(args.url, args.base, args.usuario, args.clave, verificar_ssl=not args.sin_verificar_ssl)
    print("Conectado a", odoo.url, "como", args.usuario)
    inicio = time.time()
    pasos = list(PASOS) if args.paso == "todo" else [args.paso]
    for paso in pasos:
        PASOS[paso](odoo, args)
    print("\nTerminado en %.1f minutos." % ((time.time() - inicio) / 60))


if __name__ == "__main__":
    main()
