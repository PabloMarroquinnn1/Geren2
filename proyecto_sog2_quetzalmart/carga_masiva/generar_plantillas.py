"""Genera las plantillas de carga masiva (Excel) a partir de catalogo.py.

Cada archivo se importa en Odoo con  Lista -> (engranaje) -> Importar registros.
Los encabezados son nombres técnicos de los campos, así Odoo relaciona cada
columna sin tener que elegirla a mano, en cualquier idioma de la interfaz.

Uso:  python carga_masiva/generar_plantillas.py
"""
import random
import re
import unicodedata
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import catalogo

SALIDA = Path(__file__).resolve().parent / "plantillas"
azar = random.Random(2026)

PAISES = {"gt": "base.gt", "mx": "base.mx", "sv": "base.sv"}
TELEFONO = {"gt": "+502 {} {}", "mx": "+52 962 {} {}", "sv": "+503 {} {}"}


def slug(texto):
    texto = unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "_", texto.lower()).strip("_")


def telefono(pais):
    return TELEFONO[pais].format(azar.randint(2200, 7999), azar.randint(1000, 9999))


def ean13(base12):
    """Código de barras EAN-13 válido (prefijo 740 = Guatemala)."""
    digitos = [int(d) for d in base12]
    suma = sum(d * (3 if i % 2 else 1) for i, d in enumerate(digitos))
    return base12 + str((10 - suma % 10) % 10)


def guardar(nombre, encabezados, filas, ayuda):
    libro = Workbook()
    hoja = libro.active
    hoja.title = nombre.split("_", 1)[1][:31]
    hoja.append(encabezados)
    for fila in filas:
        hoja.append(fila)
    relleno = PatternFill("solid", fgColor="0B6E4F")
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = relleno
        celda.alignment = Alignment(horizontal="center")
    for indice, encabezado in enumerate(encabezados, 1):
        largo = max([len(str(encabezado))] + [len(str(f[indice - 1] or "")) for f in filas])
        hoja.column_dimensions[get_column_letter(indice)].width = min(max(12, largo + 2), 60)
    hoja.freeze_panes = "A2"
    notas = libro.create_sheet("instrucciones")
    notas.append(["Cómo importar este archivo en Odoo"])
    notas["A1"].font = Font(bold=True, size=13, color="0B6E4F")
    for linea in ayuda:
        notas.append([linea])
    notas.column_dimensions["A"].width = 110
    ruta = SALIDA / f"{nombre}.xlsx"
    libro.save(ruta)
    print(f"  {ruta.name:<34} {len(filas):>4} filas")


# ---------------------------------------------------------------------------
def departamentos():
    filas = [[f"qm_dep_{codigo}", nombre] for codigo, nombre in catalogo.DEPARTAMENTOS]
    guardar("01_departamentos", ["id", "name"], filas, [
        "Empleados > Departamentos > (engranaje) > Importar registros > Subir archivo.",
        "La columna id es el identificador externo: si vuelve a importar, actualiza en lugar de duplicar.",
    ])


def puestos():
    filas = [[f"qm_puesto_{codigo}", nombre, f"qm_dep_{depto}", 2,
              f"Puesto de {nombre.lower()} en las sucursales de QuetzalMart."]
             for codigo, nombre, depto, _salario in catalogo.PUESTOS]
    guardar("02_puestos", ["id", "name", "department_id/id", "no_of_recruitment", "description"], filas, [
        "Empleados > Configuración > Puestos de trabajo > Importar registros.",
        "Importe primero 01_departamentos: department_id/id hace referencia a esos identificadores.",
    ])


def empleados():
    gerentes = {}
    filas = []
    orden = sorted(enumerate(catalogo.EMPLEADOS, 1),
                   key=lambda par: {"director": 0, "gerente_sucursal": 1}.get(par[1][1], 2))
    for numero, (nombre, puesto, sucursal, genero) in orden:
        partes = nombre.split()
        correo = f"{slug(partes[0])}.{slug(partes[-2])}@quetzalmart.example"
        datos_puesto = next(p for p in catalogo.PUESTOS if p[0] == puesto)
        if puesto == "director":
            jefe = ""
        elif puesto == "gerente_sucursal":
            jefe = "qm_emp_001"
            gerentes[sucursal] = f"qm_emp_{numero:03d}"
        else:
            jefe = gerentes.get(sucursal, "qm_emp_001")
        pais = {"GUA": "gt", "MEX": "mx", "SAL": "sv"}[sucursal]
        nacimiento = date(1975, 1, 1) + timedelta(days=azar.randint(0, 28 * 365))
        filas.append([
            f"qm_emp_{numero:03d}", nombre, f"qm_puesto_{puesto}", datos_puesto[1],
            f"qm_dep_{datos_puesto[2]}", f"qm_implementacion.ubicacion_trabajo_{sucursal.lower()}",
            jefe, correo, telefono(pais), genero, nacimiento.isoformat(),
            "".join(str(azar.randint(0, 9)) for _ in range(13)), PAISES[pais],
            f"{slug(partes[0])}{azar.randint(10, 99)}@correo.example", "employee",
        ])
    guardar("03_empleados", [
        "id", "name", "job_id/id", "job_title", "department_id/id", "work_location_id/id",
        "parent_id/id", "work_email", "mobile_phone", "gender", "birthday", "identification_id",
        "private_country_id/id", "private_email", "employee_type",
    ], filas, [
        "Empleados > (engranaje) > Importar registros.",
        "Los jefes (parent_id) aparecen primero en el archivo para que existan antes que su equipo.",
        "identification_id es el número de DPI / identificación del empleado.",
    ])


def contratos():
    filas = []
    for numero, (nombre, puesto, _sucursal, _genero) in enumerate(catalogo.EMPLEADOS, 1):
        _codigo, nombre_puesto, depto, salario = next(p for p in catalogo.PUESTOS if p[0] == puesto)
        inicio = date(2024, 1, 8) + timedelta(days=azar.randint(0, 600))
        filas.append([
            f"qm_contrato_{numero:03d}", f"Contrato {nombre_puesto} - {nombre}",
            f"qm_emp_{numero:03d}", f"qm_puesto_{puesto}", f"qm_dep_{depto}",
            inicio.isoformat(), round(salario * azar.uniform(0.95, 1.12), 2), "open",
        ])
    guardar("04_contratos", ["id", "name", "employee_id/id", "job_id/id", "department_id/id",
                             "date_start", "wage", "state"], filas, [
        "Empleados > Contratos > Importar registros (requiere haber importado 03_empleados).",
        "state = open significa contrato En proceso (vigente).",
    ])


def clientes():
    filas = []
    for numero, (nombre, tipo, pais, ciudad, estado, nit) in enumerate(catalogo.CLIENTES, 1):
        etiqueta = "Cliente empresarial" if tipo == "company" else "Cliente consumidor final"
        correo = f"compras.{slug(nombre)[:24]}@correo.example" if tipo == "company" \
            else f"{slug(nombre.split()[0])}.{slug(nombre.split()[-2])}@correo.example"
        filas.append([
            f"qm_cliente_{numero:03d}", nombre, tipo, nit, correo, telefono(pais),
            f"{azar.randint(1, 40)}a. Calle {azar.randint(1, 30)}-{azar.randint(10, 90)}",
            ciudad, estado or "", PAISES[pais], f"CLI-{numero:03d}", etiqueta, 1, "es_419",
        ])
    guardar("05_clientes", ["id", "name", "company_type", "vat", "email", "phone", "street", "city",
                            "state_id", "country_id/id", "ref", "category_id", "customer_rank", "lang"],
            filas, [
                "Contactos > (engranaje) > Importar registros.",
                "En la prueba de importación marque 'Crear si no existe' para la columna category_id (etiquetas).",
            ])


def proveedores():
    filas = []
    for numero, (nombre, pais, ciudad, nit, _rubro) in enumerate(catalogo.PROVEEDORES, 1):
        filas.append([
            f"qm_proveedor_{numero:03d}", nombre, "company", nit,
            f"ventas.{slug(nombre)[:22]}@proveedor.example", telefono(pais),
            f"Zona industrial, bodega {azar.randint(1, 60)}", ciudad, PAISES[pais],
            f"PRV-{numero:03d}", "Proveedor", 1, "es_419",
        ])
    for numero, (nombre, servicio, _monto) in enumerate(catalogo.SERVICIOS_OUTSOURCING, 1):
        filas.append([
            f"qm_outsourcing_{numero:03d}", nombre, "company",
            f"{azar.randint(1000000, 9999999)}-{azar.randint(0, 9)}",
            f"contratos.{slug(nombre)[:20]}@proveedor.example", telefono("gt"),
            f"Oficinas centrales, nivel {azar.randint(2, 12)}", "Ciudad de Guatemala", PAISES["gt"],
            f"OUT-{numero:03d}", "Outsourcing", 1, "es_419",
        ])
    guardar("06_proveedores", ["id", "name", "company_type", "vat", "email", "phone", "street", "city",
                               "country_id/id", "ref", "category_id", "supplier_rank", "lang"], filas, [
        "Contactos > Importar registros. supplier_rank = 1 los marca como proveedores.",
        "Incluye los 5 proveedores de servicios tercerizados (contratos de outsourcing del gestor documental).",
    ])


def proveedor_de(rubro, indice):
    candidatos = [f"qm_proveedor_{n:03d}" for n, p in enumerate(catalogo.PROVEEDORES, 1) if p[4] == rubro]
    return candidatos[indice % len(candidatos)]


def productos():
    filas = []
    for orden_cat, (clave, datos) in enumerate(catalogo.CATEGORIAS_VENTA.items()):
        for numero, (nombre, presentacion, costo, precio, peso, _envase, _color) in enumerate(datos["productos"], 1):
            sku = f"{datos['prefijo']}-{numero:03d}"
            filas.append([
                f"qm_producto_{slug(sku)}", f"{nombre} {presentacion}", sku,
                ean13(f"740{orden_cat + 1:02d}{numero:07d}"), f"qm_implementacion.cat_{clave}",
                f"qm_implementacion.web_{clave}", "consu", 1, precio, costo, peso,
                f"{nombre} {presentacion}. {datos['descripcion']}",
                1, 1, 1, proveedor_de(clave, numero), round(costo, 2),
            ])
    guardar("07_productos", [
        "id", "name", "default_code", "barcode", "categ_id/id", "public_categ_ids/id", "type",
        "is_storable", "list_price", "standard_price", "weight", "description_sale",
        "is_published", "sale_ok", "purchase_ok", "seller_ids/partner_id/id", "seller_ids/price",
    ], filas, [
        "Inventario o Ventas > Productos > Importar registros.",
        "list_price es el precio de venta con IVA incluido (el IVA de Guatemala, 12 %, se desglosa en la factura).",
        "seller_ids define el proveedor principal y su precio de compra.",
        "Las imágenes se cargan después con: python carga_masiva/cargar_quetzalmart.py --paso imagenes",
    ])


def materiales():
    filas = []
    numero = 0
    for clave, (_nombre_cat, items) in catalogo.CATEGORIAS_MATERIALES.items():
        for indice, (nombre, unidad, costo, _stock) in enumerate(items):
            numero += 1
            sku = f"MAT-{numero:03d}"
            filas.append([
                f"qm_material_{numero:03d}", nombre, sku, f"qm_implementacion.cat_mat_{clave}",
                "consu", 1, 0, 1, 0, costo, f"Unidad de compra: {unidad}. Uso interno de las sucursales.",
                proveedor_de("materiales", indice), costo,
            ])
    guardar("08_materiales_operativos", [
        "id", "name", "default_code", "categ_id/id", "type", "is_storable", "sale_ok", "purchase_ok",
        "is_published", "standard_price", "description_purchase", "seller_ids/partner_id/id",
        "seller_ids/price",
    ], filas, [
        "Inventario > Productos > Importar registros.",
        "Son insumos para que cada sucursal funcione: no se venden (sale_ok = 0) ni se publican en la web.",
    ])


def existencias():
    sector_de = {clave: datos["sector"] for clave, datos in catalogo.CATEGORIAS_VENTA.items()}
    filas = []
    for clave, datos in catalogo.CATEGORIAS_VENTA.items():
        for numero in range(1, len(datos["productos"]) + 1):
            sku = f"{datos['prefijo']}-{numero:03d}"
            for sucursal, base in (("gua", 120), ("mex", 80), ("sal", 70)):
                cantidad = base + azar.randint(0, 60)
                filas.append([sku, f"qm_implementacion.ubicacion_{sucursal}_{sector_de[clave]}", cantidad])
    numero = 0
    for _clave, (_nombre, items) in catalogo.CATEGORIAS_MATERIALES.items():
        for (_n, _u, _c, stock) in items:
            numero += 1
            for sucursal in ("gua", "mex", "sal"):
                filas.append([f"MAT-{numero:03d}", f"qm_implementacion.ubicacion_{sucursal}_materiales", stock])
    guardar("09_existencias_iniciales", ["product_id", "location_id/id", "inventory_quantity"], filas, [
        "Inventario > Operaciones > Inventario físico > Importar registros.",
        "Después de importar, seleccione todo y pulse 'Aplicar' para registrar las existencias.",
        "Cada producto queda en el sector que le corresponde de cada sucursal (secos, refrigeración, etc.).",
    ])


def oportunidades():
    etapas = ["crm.stage_lead1", "crm.stage_lead2", "crm.stage_lead3", "crm.stage_lead4"]
    filas = []
    empresas = [(n, c) for n, c in enumerate(catalogo.CLIENTES, 1) if c[1] == "company"]
    for indice, (numero, (nombre, _t, pais, *_r)) in enumerate(empresas[:18]):
        equipo = {"gt": "sales_team.team_sales_department", "mx": "qm_implementacion.equipo_mex",
                  "sv": "qm_implementacion.equipo_sal"}[pais]
        etapa = etapas[indice % 4]
        filas.append([
            f"qm_oportunidad_{indice + 1:03d}",
            f"Abastecimiento mensual - {nombre}", "opportunity", f"qm_cliente_{numero:03d}",
            equipo, etapa, azar.choice([2500, 4800, 7200, 12500, 18000, 26000]),
            {"crm.stage_lead1": 10, "crm.stage_lead2": 35, "crm.stage_lead3": 70, "crm.stage_lead4": 100}[etapa],
            (date(2026, 11, 1) + timedelta(days=azar.randint(0, 75))).isoformat(),
            "Cliente empresarial interesado en un contrato de abastecimiento con entrega programada.",
        ])
    guardar("10_oportunidades_crm", ["id", "name", "type", "partner_id/id", "team_id/id", "stage_id/id",
                                     "expected_revenue", "probability", "date_deadline", "description"],
            filas, ["CRM > Ventas > Mi flujo > (engranaje) > Importar registros."])


def main():
    SALIDA.mkdir(exist_ok=True)
    print("Plantillas de carga masiva:")
    departamentos()
    puestos()
    empleados()
    contratos()
    clientes()
    proveedores()
    productos()
    materiales()
    existencias()
    oportunidades()


if __name__ == "__main__":
    main()
